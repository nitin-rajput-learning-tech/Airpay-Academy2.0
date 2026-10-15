"""Convert a C-suite brief written in Markdown to .docx in the Sentientia Product Guide format.

The format mirrors .tools/docgen/build-product-guide.js: A4, 1 inch margins, Arial 10.5 pt, brand blue 0066A7,
heading 1 with a blue bottom rule, heading 2 in brand blue, tinted table header row (E8F2F9) with grey borders,
a footer with the title and "Page X of Y". The Markdown file stays the single source of truth; the .docx is
generated from it so the two cannot drift.

Supported Markdown subset: one H1 (title block), a metadata paragraph right after it, H2/H3, paragraphs with
**bold**, *italic* and `code`, "-" bullet lists, "1." numbered lists, pipe tables, "---" rules (ignored) and a
trailing italic note.

Usage:  python -I md_to_docx.py <input.md> <output.docx> "<footer title>"
Requires python-docx (pip install python-docx).
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor

BRAND = RGBColor(0x00, 0x66, 0xA7)
TXT = RGBColor(0x1A, 0x1A, 0x2E)
GREY = RGBColor(0x5A, 0x60, 0x70)
TINT = "E8F2F9"
BORDER = "C9D2E0"
CONTENT_DXA = 9026
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def set_cell_shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tbl_pr.append(borders)


def set_cell_margins(table) -> None:
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", 80), ("left", 120), ("bottom", 80), ("right", 120)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def paragraph_border(paragraph, side: str, size: int, color: str, space: int = 4) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bdr = p_pr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        p_pr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    bdr.append(el)


def add_runs(paragraph, text: str, *, size: float | None = None, color: RGBColor | None = TXT, bold: bool = False, italic: bool = False) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        run_bold, run_italic, run_code, body = bold, italic, False, part
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run_bold, body = True, part[2:-2]
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run_code, body = True, part[1:-1]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_italic, body = True, part[1:-1]
        run = paragraph.add_run(body)
        run.bold = run_bold
        run.italic = run_italic
        if run_code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


def add_field(paragraph, instr: str, size: float, color: RGBColor) -> None:
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = color
    for kind, text in (("begin", None), ("instr", instr), ("end", None)):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def column_widths(rows: list[list[str]]) -> list[int]:
    cols = len(rows[0])
    weights = []
    for c in range(cols):
        longest = max(len(re.sub(r"[*`]", "", r[c])) for r in rows if c < len(r))
        weights.append(min(max(longest, 10), 70) ** 0.8)
    total = sum(weights)
    widths = [max(900, int(CONTENT_DXA * w / total)) for w in weights]
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_table(doc, rows: list[list[str]]) -> None:
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    widths = column_widths(rows)
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.width = Emu(widths[c_idx] * 635)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            add_runs(para, text, size=9.5, bold=(r_idx == 0))
            if r_idx == 0:
                set_cell_shade(cell, TINT)
    # repeat the header row across pages
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def build(md_text: str, footer_title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Emu(11906 * 635), Emu(16838 * 635)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Emu(1440 * 635))
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TXT
    for name, size, color in (("Heading 1", 15, TXT), ("Heading 2", 12, BRAND)):
        st = doc.styles[name]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st.element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        for theme_attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            st.element.rPr.rFonts.attrib.pop(qn(theme_attr), None)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
    # footer
    foot = section.footer.paragraphs[0]
    add_runs(foot, f"{footer_title} · Confidential · ", size=8, color=GREY)
    add_runs(foot, "Page ", size=8, color=GREY)
    add_field(foot, "PAGE", 8, GREY)
    add_runs(foot, " of ", size=8, color=GREY)
    add_field(foot, "NUMPAGES", 8, GREY)

    lines = md_text.splitlines()
    i = 0
    title_done = False
    last_para = None
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or re.fullmatch(r"-{3,}", line.strip()):
            i += 1
            continue
        if line.startswith("# ") and not title_done:
            title = line[2:].strip()
            main, _, sub = title.partition("—")
            p1 = doc.add_paragraph()
            p1.paragraph_format.space_before = Pt(30)
            p1.paragraph_format.space_after = Pt(3)
            add_runs(p1, main.strip().upper() if "Sentientia" in main else main.strip(), size=28, color=BRAND, bold=True)
            p2 = doc.add_paragraph()
            p2.paragraph_format.space_after = Pt(4)
            add_runs(p2, sub.strip() or "Brief", size=20, color=TXT)
            title_done = True
            i += 1
            # metadata paragraph
            while i < len(lines) and not lines[i].strip():
                i += 1
            if i < len(lines) and lines[i].startswith("**"):
                meta = doc.add_paragraph()
                meta.paragraph_format.space_after = Pt(24)
                paragraph_border(meta, "bottom", 8, "0066A7", 6)
                add_runs(meta, lines[i].strip().replace("**", ""), size=9.5, color=GREY)
                i += 1
            continue
        if line.startswith("## "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(9)
            paragraph_border(h, "bottom", 6, "0066A7", 2)
            add_runs(h, line[3:].strip(), size=15, color=TXT, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(11)
            h.paragraph_format.space_after = Pt(6)
            add_runs(h, line[4:].strip(), size=12, color=BRAND, bold=True)
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        m = re.match(r"^(\s*)(-|\d+\.)\s+(.*)$", line)
        if m:
            numbered = m.group(2)[0].isdigit()
            text = m.group(3)
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not re.match(r"^\s*(-|\d+\.)\s", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            if numbered:
                # manual numbers: Word's List Number style would continue the count across separate lists
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(20)
                p.paragraph_format.first_line_indent = Pt(-20)
                p.paragraph_format.tab_stops.add_tab_stop(Pt(20))
                p.paragraph_format.space_after = Pt(5)
                add_runs(p, m.group(2) + "	" + text)
            else:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(5)
                add_runs(p, text)
            continue
        # paragraph (merge wrapped lines)
        text = line.strip()
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|\||-{3,}|\s*(-|\d+\.)\s)", lines[i]):
            text += " " + lines[i].strip()
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        if text.startswith("*") and text.endswith("*") and not text.startswith("**"):
            paragraph_border(p, "top", 4, BORDER, 4)
            p.paragraph_format.space_before = Pt(10)
            add_runs(p, text.strip("*"), size=8.5, color=GREY, italic=True)
        else:
            add_runs(p, text)
        last_para = p
    return doc


def main(src: str, out: str, footer_title: str) -> int:
    doc = build(Path(src).read_text(encoding="utf-8"), footer_title)
    doc.core_properties.title = footer_title
    doc.core_properties.author = "Airpay L&D"
    doc.save(out)
    print(f"Wrote {out}")
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(*sys.argv[1:]))
