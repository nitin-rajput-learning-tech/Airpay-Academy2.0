#!/usr/bin/env python3
"""md_to_docx.py - regenerate a .docx handout from one of the UAT Markdown docs.

Minimal, dependency-light (python-docx only) converter for the subset of Markdown the
cutover docs use: #/##/### headings, paragraphs, - bullets, 1. numbered steps, pipe
tables, **bold**, `code`, and --- rules. Nothing else. Written 2026-09-08 because the
team handout UAT-TEST-ACCOUNTS.docx had drifted from its Markdown source.

Usage:
    python tools/uat/md_to_docx.py <input.md> <output.docx>

The uat-credentials/ folder is gitignored (it holds passwords) - run this locally and
share the .docx by the same channel as the .md.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str) -> None:
    """Append text to a paragraph, honouring **bold** and `code` spans."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
        else:
            paragraph.add_run(part)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(src: Path, dst: Path) -> None:
    doc = Document()
    lines = src.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped == "---":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if m:
            doc.add_heading(m.group(2).replace("`", ""), level=len(m.group(1)))
            i += 1
            continue
        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not re.match(r"^\|\s*:?-{2,}", lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                table = doc.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c in range(width):
                        cell = table.cell(r, c)
                        cell.text = ""
                        add_inline(cell.paragraphs[0], row[c] if c < len(row) else "")
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
                doc.add_paragraph()
            continue
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
            i += 1
            continue
        m = re.match(r"^\d+\.\s+(.*)$", stripped)
        if m:
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
            i += 1
            continue
        # Paragraph: merge soft-wrapped lines until a blank line or a block start.
        buf = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", "- ", "* ", "---")) or re.match(r"^\d+\.\s", nxt):
                break
            buf.append(nxt)
            i += 1
        add_inline(doc.add_paragraph(), " ".join(buf))
    doc.save(dst)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
    print(f"wrote {sys.argv[2]}")
